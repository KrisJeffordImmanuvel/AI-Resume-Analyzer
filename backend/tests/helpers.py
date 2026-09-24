"""Build real PDF/DOCX files in memory so tests need no binary fixtures."""

import io


def make_pdf(lines: list[str]) -> bytes:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    y = 72
    for line in lines:
        page.insert_text((72, y), line, fontsize=11)
        y += 16
    data = doc.tobytes()
    doc.close()
    return data


def make_blank_pdf() -> bytes:
    import pymupdf

    doc = pymupdf.open()
    doc.new_page()
    data = doc.tobytes()
    doc.close()
    return data


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
