"""Text extraction for resumes (PDF, DOCX, TXT) and job descriptions (TXT or pasted text).

All extracted text goes through normalize_text(), and that normalized text is the
single source of truth: every evidence quote shown to the user is a substring of it.
"""

import io
import re
import unicodedata
from pathlib import PurePath

MAX_UPLOAD_BYTES = 5 * 1024 * 1024  # 5 MB
MAX_TEXT_CHARS = 100_000

RESUME_EXTENSIONS = {".pdf", ".docx", ".txt"}
JD_EXTENSIONS = {".txt"}


class ParseError(Exception):
    """A file could not be turned into usable text. `status` is the HTTP code to return."""

    def __init__(self, message: str, status: int = 422):
        super().__init__(message)
        self.message = message
        self.status = status


def normalize_text(text: str) -> str:
    """Unicode-normalize and tidy whitespace while keeping line structure."""
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # Bullets and odd characters PDFs often emit become plain bullets/spaces.
    text = text.replace(" ", " ").replace("​", "").replace("", "•")
    lines = [re.sub(r"[ \t\f\v]+", " ", line).strip() for line in text.split("\n")]
    text = "\n".join(lines)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _extension(filename: str) -> str:
    return PurePath(filename or "").suffix.lower()


def _check_size(data: bytes, label: str) -> None:
    if len(data) > MAX_UPLOAD_BYTES:
        raise ParseError(f"{label} is larger than {MAX_UPLOAD_BYTES // (1024 * 1024)} MB.", status=413)
    if not data:
        raise ParseError(f"{label} is empty.")


def _decode_txt(data: bytes) -> str:
    # utf-8-sig strips a BOM; utf-16 covers files saved as "Unicode" by Notepad;
    # cp1252 is the classic Windows ANSI encoding.
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _extract_pdf(data: bytes) -> str:
    import pymupdf

    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise ParseError("The PDF could not be opened. It may be damaged or not a real PDF.") from exc
    with doc:
        if doc.needs_pass:
            raise ParseError("The PDF is password-protected. Please upload an unlocked copy.")
        return "\n".join(page.get_text("text") for page in doc)


def _extract_docx(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ParseError("The DOCX file could not be opened. It may be damaged or not a real .docx.") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                if cell.text not in cells:  # merged cells repeat their text
                    cells.append(cell.text)
            parts.append(" | ".join(cells))
    return "\n".join(parts)


def _finish(raw: str, label: str, empty_hint: str = "") -> str:
    text = normalize_text(raw)
    if not text:
        raise ParseError(f"No text could be extracted from the {label}.{empty_hint}")
    if len(text) > MAX_TEXT_CHARS:
        raise ParseError(f"The {label} is too long (over {MAX_TEXT_CHARS:,} characters).", status=413)
    return text


def extract_resume_text(filename: str, data: bytes) -> str:
    ext = _extension(filename)
    if ext == ".doc":
        raise ParseError("Old .doc files are not supported. Save the resume as .docx or .pdf and try again.", 415)
    if ext not in RESUME_EXTENSIONS:
        raise ParseError("Resume must be a PDF, DOCX or TXT file.", status=415)
    _check_size(data, "Resume file")
    if ext == ".pdf":
        raw = _extract_pdf(data)
        hint = " It may be a scanned image; export it as a text-based PDF or DOCX."
    elif ext == ".docx":
        raw, hint = _extract_docx(data), ""
    else:
        raw, hint = _decode_txt(data), ""
    return _finish(raw, "resume", hint)


def extract_jd_file_text(filename: str, data: bytes) -> str:
    if _extension(filename) not in JD_EXTENSIONS:
        raise ParseError("Job description files must be .txt. You can also paste the text instead.", status=415)
    _check_size(data, "Job description file")
    return _finish(_decode_txt(data), "job description")


def clean_jd_text(text: str) -> str:
    return _finish(text or "", "job description")
